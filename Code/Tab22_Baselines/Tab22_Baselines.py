"""Tab 22 - baselines (seasonal-naive + ARIMA), task 10.

Reads Code/experiments.csv filtered to tab_name == "Tab22_Baselines", loops every row,
records point forecasts + runtime for BOTH baseline methods per origin (one sheet per
exp_id, same as every other tab). No AutoGluon, no GPU, no ensemble weights - these fit
in seconds.

Run: python Code/Tab22_Baselines/Tab22_Baselines.py [--exp-id T22-M-001] [--force]
"""
from __future__ import annotations

import argparse
import csv
import sys
import time
from pathlib import Path

import pandas as pd

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT))

from Code.common.baselines import run_baselines  # noqa: E402
from Code.common.data import FREQ_ALIAS, FREQ_TO_SEASONAL_PERIOD, build_scoring_tsdf, build_train_tsdf, load_final_data, parse_plan_date  # noqa: E402
from Code.common.results_io import RESULTS_DIR, write_block_header, write_dataframe_block, already_has_result, write_run_config  # noqa: E402

import openpyxl  # noqa: E402

TAB_NAME = "Tab22_Baselines"
EXPERIMENTS_CSV = REPO_ROOT / "Code" / "experiments.csv"


def load_tab_rows() -> list[dict]:
    with open(EXPERIMENTS_CSV, encoding="utf-8") as f:
        rows = [r for r in csv.DictReader(f) if r["tab_name"] == TAB_NAME]
    if not rows:
        raise RuntimeError(
            f"No rows for {TAB_NAME} in {EXPERIMENTS_CSV}. Run "
            "Code/common/update_plan_for_local.py then Code/common/regenerate_configs.py first."
        )
    return rows


def run_one(exp_row: dict) -> None:
    index = exp_row["index"]
    frequency = exp_row["frequency"]
    freq_key = FREQ_ALIAS[frequency]
    seasonal_period = FREQ_TO_SEASONAL_PERIOD[freq_key]

    train_data_end = parse_plan_date(exp_row["train_data_end"])
    horizon_end = parse_plan_date(exp_row["horizon_end"])
    prediction_length = int(exp_row["n_periods"])

    series = load_final_data(index, freq_key)
    train_tsdf = build_train_tsdf(series, train_data_end, item_id=index)
    scoring_tsdf = build_scoring_tsdf(series, train_data_end, horizon_end, item_id=index)

    past_values = train_tsdf["target"].to_numpy()
    actual_h = scoring_tsdf["target"].to_numpy()[-prediction_length:]
    horizon_dates = scoring_tsdf.index.get_level_values("timestamp")[-prediction_length:]

    t0 = time.monotonic()
    out = run_baselines(past_values, prediction_length, seasonal_period)
    total_runtime = time.monotonic() - t0

    forecast_df = pd.DataFrame({
        "date": horizon_dates,
        "actual": actual_h,
        "seasonal_naive": out["seasonal_naive_forecast"],
        "arima": out["arima_forecast"],
    })

    metadata = {
        "exp_id": exp_row["exp_id"], "tab_name": TAB_NAME, "index": index, "frequency": frequency,
        "train_data_end": train_data_end.date().isoformat(), "horizon_end": horizon_end.date().isoformat(),
        "n_periods": prediction_length,
        "seasonal_naive_runtime_sec": round(out["seasonal_naive_runtime_sec"], 4),
        "arima_runtime_sec": round(out["arima_runtime_sec"], 4),
        "arima_order": str(out["arima_order"]),
        "total_runtime_sec": round(total_runtime, 4),
    }

    _write_sheet(exp_row["exp_id"], metadata, forecast_df)
    write_run_config(TAB_NAME, exp_row["exp_id"], {**metadata, "method": "seasonal_naive + ARIMA(grid search by AIC)"})
    print(f"  {exp_row['exp_id']}: done in {total_runtime:.2f}s (ARIMA order {out['arima_order']})")


def _write_sheet(exp_id: str, metadata: dict, forecast_df: pd.DataFrame) -> None:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    workbook_path = RESULTS_DIR / f"Results_{TAB_NAME}.xlsx"
    if workbook_path.exists():
        wb = openpyxl.load_workbook(workbook_path)
        if exp_id in wb.sheetnames:
            del wb[exp_id]
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
    ws = wb.create_sheet(title=exp_id)
    row = write_block_header(ws, 1, f"Experiment {exp_id} (baseline: seasonal-naive + ARIMA)")
    meta_df = pd.DataFrame({"field": list(metadata.keys()), "value": list(metadata.values())})
    row = write_dataframe_block(ws, row, meta_df)
    row = write_block_header(ws, row, "Forecast (both baseline methods; no ensemble, no weights)")
    write_dataframe_block(ws, row, forecast_df)
    wb.save(workbook_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--exp-id", default=None)
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()

    rows = load_tab_rows()
    if args.exp_id:
        rows = [r for r in rows if r["exp_id"] == args.exp_id]
        if not rows:
            raise SystemExit(f"exp_id {args.exp_id!r} not found in {TAB_NAME}")

    workbook_path = RESULTS_DIR / f"Results_{TAB_NAME}.xlsx"
    n_skipped = 0
    for row in rows:
        if not args.force and already_has_result(workbook_path, row["exp_id"]):
            n_skipped += 1
            continue
        run_one(row)
    print(f"\n{TAB_NAME}: {len(rows) - n_skipped} run, {n_skipped} skipped (already had results)")


if __name__ == "__main__":
    main()
