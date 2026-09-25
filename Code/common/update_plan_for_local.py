"""Task 9 + task 10: update Main_Experiment_Plan.xlsx for local execution.

  1. Rename the 'notebook' header -> 'script' (MASTER + every Tab01..Tab21 sheet), and
     rewrite every .ipynb path to .py. Values only - cell.fill/font are never touched,
     so colour coding survives untouched.
  2. Fill CONTROLS with what is actually used (see Code/common/experiment.py, the
     single source of truth these values are read from).
  3. Add Tab 22 (baselines): 294 new rows in MASTER, derived from Tabs 1-3's own rows
     so the origins are identical by construction (not re-derived and risking drift),
     plus a new Tab22_Baselines sheet following the workbook's own conventions.

This script does NOT run regenerate_configs.py itself - run that separately afterward,
so a partial/failed regenerate never leaves the workbook edit half-done.
"""
from __future__ import annotations

import copy
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT))

from Code.common.data import parse_plan_date  # noqa: E402
from Code.common.experiment import (  # noqa: E402
    NUM_VAL_WINDOWS,
    QUANTILE_LEVELS,
    RANDOM_SEED,
    TIME_LIMIT_S,
    build_hyperparameters,
)
from Code.common import gpu as gpu_module  # noqa: E402

PLAN_PATH = REPO_ROOT / "Main_Experiment_Plan.xlsx"

GREEN = "FFE2EFDA"
YELLOW = "FFFFF2CC"
HEADER_FILL_RGB = "FF1F3864"

# ---------------------------------------------------------------------------
# 1. notebook -> script
# ---------------------------------------------------------------------------

def rename_notebook_to_script(wb: openpyxl.Workbook) -> None:
    for sheet in wb.sheetnames:
        if sheet != "MASTER" and not sheet.startswith("Tab"):
            continue
        ws = wb[sheet]
        header_row = 4
        notebook_col = None
        for c in range(1, ws.max_column + 1):
            if ws.cell(row=header_row, column=c).value == "notebook":
                notebook_col = c
                break
        if notebook_col is None:
            continue
        ws.cell(row=header_row, column=notebook_col, value="script")
        for r in range(header_row + 1, ws.max_row + 1):
            cell = ws.cell(row=r, column=notebook_col)
            if isinstance(cell.value, str) and cell.value.endswith(".ipynb"):
                cell.value = cell.value[: -len(".ipynb")] + ".py"


# ---------------------------------------------------------------------------
# 2. CONTROLS
# ---------------------------------------------------------------------------

def fill_controls(wb: openpyxl.Workbook) -> None:
    import autogluon.timeseries as agts

    try:
        gpu_info = gpu_module.require_cuda()
        hardware_value = f"{gpu_info.name} (local GPU, recorded per run; VRAM {gpu_info.vram_gb:.1f} GB)"
    except RuntimeError as exc:
        hardware_value = f"UNVERIFIED AT PLAN-UPDATE TIME - {exc}"

    def _describe(hp: dict) -> str:
        return ", ".join(f"{k}[{len(v)} config(s)]" if isinstance(v, list) else k for k, v in hp.items())

    model_pool_str = (
        f"TWO POOLS, fixed: Sets 1 & 2 = explicit hyperparameters dict with pretrained "
        f"foundation models EXCLUDED (see Code/common/experiment.py:build_hyperparameters): "
        f"{_describe(build_hyperparameters(include_pretrained=False))}. Set 3 = the same "
        f"pool WITH them included: {_describe(build_hyperparameters(include_pretrained=True))}."
    )

    values = {
        "Framework version": (agts.__version__, True),
        "Preset / model pool": (model_pool_str, True),
        "time_limit per fit (seconds)": (str(TIME_LIMIT_S), False),
        "num_val_windows": (str(NUM_VAL_WINDOWS), True),
        "Quantile levels": (", ".join(str(q) for q in QUANTILE_LEVELS), True),
        "Target transformation": ("Index level, 2021=100 (frozen decision, P1-05)", True),
        "Random seed": (str(RANDOM_SEED), True),
        "Data source and vintage": ("Data/final_data.csv - see Data/final_data_notes.md", True),
        "Series start date": ("2014-01 (monthly/quarterly); 2014 (annual) - all three indices", True),
        "Hardware": (hardware_value, True),
    }

    ws = wb["CONTROLS"]
    for r in range(5, ws.max_row + 1):
        label = ws.cell(row=r, column=1).value
        if label in values:
            value, fixed = values[label]
            ws.cell(row=r, column=2, value=value)
            ws.cell(row=r, column=3, value="Yes" if fixed else "Provisional - see P2-09")


# ---------------------------------------------------------------------------
# 3. Tab 22 baselines
# ---------------------------------------------------------------------------

FREQ_CODE = {"Monthly": "M", "Quarterly": "Q", "Annual": "A"}


def _collect_source_rows(wb: openpyxl.Workbook) -> list[dict]:
    hdr = [wb["MASTER"].cell(row=4, column=c).value for c in range(1, 26)]
    rows = []
    for sheet in ["Tab01_HCPI_EM1", "Tab02_FCPI_EM1", "Tab03_CCPI_EM1"]:
        ws = wb[sheet]
        for r in range(5, ws.max_row + 1):
            row = {hdr[c - 1]: ws.cell(row=r, column=c).value for c in range(1, 26)}
            if row["exp_id"]:
                rows.append(row)
    return rows


def build_tab22_rows(wb: openpyxl.Workbook) -> list[dict]:
    source_rows = _collect_source_rows(wb)
    index_order = {"HCPI": 0, "FCPI": 1, "CCPI": 2}
    by_freq: dict[str, list[dict]] = {"Monthly": [], "Quarterly": [], "Annual": []}
    for row in source_rows:
        by_freq[row["frequency"]].append(row)

    new_rows = []
    for freq, freq_rows in by_freq.items():
        freq_rows_sorted = sorted(
            freq_rows,
            key=lambda r: (index_order[r["index"]], parse_plan_date(r["horizon_start"])),
        )
        for i, src in enumerate(freq_rows_sorted, start=1):
            exp_id = f"T22-{FREQ_CODE[freq]}-{i:03d}"
            new_rows.append({
                "exp_id": exp_id,
                "tab": "Tab 22",
                "tab_name": "Tab22_Baselines",
                "index": src["index"],
                "index_name": src["index_name"],
                "eval_metric": "N/A",
                "metric_name": "Baseline (no fitting objective)",
                "metric_family": "Baseline",
                "frequency": freq,
                "horizon_start": src["horizon_start"],
                "horizon_end": src["horizon_end"],
                "n_periods": src["n_periods"],
                "train_data_end": src["train_data_end"],
                "train_data_start": src["train_data_start"],
                "regime_of_horizon_start": src["regime_of_horizon_start"],
                "records": "Point + runtime",
                "script": "Code/Tab22_Baselines/Tab22_Baselines.py",
                "tab_config": "Code/common/config_Tab22_Baselines.yaml",
                "result_workbook": "Results/Results_Tab22_Baselines.xlsx",
                "result_tab": exp_id,
                "status": "Not started",
                "runtime_sec": None,
                "hardware": None,
                "run_date": None,
                "notes": None,
            })
    return new_rows


def _copy_row_style(ws, from_row: int, to_row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        src_cell = ws.cell(row=from_row, column=c)
        dst_cell = ws.cell(row=to_row, column=c)
        if src_cell.has_style:
            dst_cell.fill = copy.copy(src_cell.fill)
            dst_cell.font = copy.copy(src_cell.font)
            dst_cell.border = copy.copy(src_cell.border)
            dst_cell.number_format = src_cell.number_format


def remove_existing_tab22_rows(wb: openpyxl.Workbook) -> None:
    """Idempotency guard: if this script has run before, MASTER already has a Tab 22
    block (always appended as one contiguous run at the end - see
    append_tab22_to_master). Delete it before re-appending, or re-running duplicates
    294 rows every time.
    """
    ws = wb["MASTER"]
    tab_name_col = 3  # 'tab_name' - see the MASTER header order
    first_tab22_row = None
    for r in range(5, ws.max_row + 1):
        if ws.cell(row=r, column=tab_name_col).value == "Tab22_Baselines":
            first_tab22_row = r
            break
    if first_tab22_row is not None:
        n = ws.max_row - first_tab22_row + 1
        ws.delete_rows(first_tab22_row, n)


def append_tab22_to_master(wb: openpyxl.Workbook, rows: list[dict]) -> None:
    ws = wb["MASTER"]
    hdr = [ws.cell(row=4, column=c).value for c in range(1, 26)]
    style_row = 5  # copy fills/fonts from an existing data row
    start_row = ws.max_row + 1
    for offset, row in enumerate(rows):
        r = start_row + offset
        _copy_row_style(ws, style_row, r, len(hdr))
        for c, key in enumerate(hdr, start=1):
            ws.cell(row=r, column=c, value=row.get(key))


def create_tab22_sheet(wb: openpyxl.Workbook, rows: list[dict]) -> None:
    if "Tab22_Baselines" in wb.sheetnames:
        del wb["Tab22_Baselines"]
    ws = wb.create_sheet("Tab22_Baselines")

    header_fill = PatternFill(fill_type="solid", fgColor=HEADER_FILL_RGB)
    header_font = Font(color="FFFFFFFF", bold=True)
    green_fill = PatternFill(fill_type="solid", fgColor=GREEN)
    yellow_fill = PatternFill(fill_type="solid", fgColor=YELLOW)

    ws.cell(row=1, column=1, value="Tab 22 - Baselines (seasonal-naive + ARIMA), all 3 indices, all 3 frequencies")
    ws.cell(row=2, column=1, value=(
        f"{len(rows)} rows. Records: Point + runtime. No ensemble - no weights recorded. "
        f"Runs over the IDENTICAL origins as Tabs 1-3. Not part of Stage 1 or Stage 3."
    ))

    hdr = list(rows[0].keys())
    for c, name in enumerate(hdr, start=1):
        cell = ws.cell(row=4, column=c, value=name)
        cell.fill = header_fill
        cell.font = header_font

    for r_off, row in enumerate(rows):
        r = 5 + r_off
        for c, key in enumerate(hdr, start=1):
            cell = ws.cell(row=r, column=c, value=row.get(key))
            if key == "train_data_end":
                cell.fill = green_fill
            elif key in ("status", "runtime_sec", "hardware", "run_date", "notes"):
                cell.fill = yellow_fill


# ---------------------------------------------------------------------------

def main() -> None:
    wb = openpyxl.load_workbook(PLAN_PATH)

    rename_notebook_to_script(wb)
    fill_controls(wb)

    tab22_rows = build_tab22_rows(wb)
    assert len(tab22_rows) == 294, f"expected 294 Tab 22 rows, got {len(tab22_rows)}"
    remove_existing_tab22_rows(wb)
    append_tab22_to_master(wb, tab22_rows)
    create_tab22_sheet(wb, tab22_rows)

    wb.save(PLAN_PATH)
    print(f"Updated {PLAN_PATH}")
    print(f"  Added {len(tab22_rows)} Tab 22 rows to MASTER (now {wb['MASTER'].max_row - 4} data rows)")
    print("  Created Tab22_Baselines sheet")
    print("  Filled CONTROLS")
    print("  Renamed notebook -> script across MASTER + Tab01..Tab21")
    print("\nNext: run `python Code/common/regenerate_configs.py` to regenerate experiments.csv + YAMLs.")


if __name__ == "__main__":
    main()
