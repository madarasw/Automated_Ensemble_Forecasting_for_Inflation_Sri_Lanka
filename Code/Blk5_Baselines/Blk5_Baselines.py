"""Blk5_Baselines - the baseline arm of Block 4.1 (short-window skill test).

3 indices (HCPI, FCPI, CCPI) x 69 monthly origins (o1-o69, Jan 2020 - Sep 2025 horizon
starts) = 207 runs, each paired with a Blk4_ShortWindow model row on the identical origin
and identical 72-month window:

    HCPI  BL-M-104 .. BL-M-172
    FCPI  BL-M-173 .. BL-M-241
    CCPI  BL-M-242 .. BL-M-310

Records all five benchmarks returned by Code/common/baselines.run_baselines() with
seasonal_period=12 - naive (random walk), drift, atkeson_ohanian, seasonal_naive, arima -
with each method's runtime and the selected ARIMA order. Adapted from
Code/Tab22_Baselines/Tab22_Baselines.py, with two deliberate differences: the training
series is the FIXED 72-month window [train_data_start, train_data_end] (Tab22 used an
expanding window - which would give ARIMA a different sample from the paired model row),
and the rows are selected by block rather than tab_name.

Row filter: block == "4.1" AND exp_id starts with "BL-". There are 441 BL-* rows in the
plan and 234 belong to other blocks (2.1, 2.2, 2.3, 3.1, 3.2, 3.3, 4.2); block == "4.1"
alone also returns the 207 B41-M-* model rows. validate_rows() asserts the design and every
row's window against Data/final_data.csv before anything is run.

One sheet per exp_id in Results/Results_Blk5_Baselines.xlsx, a run_config.json per run
under Results/RunConfigs/Blk5_Baselines/, resume by skipping existing sheets.

Run: python Code/Blk5_Baselines/Blk5_Baselines.py [--exp-id <id>] [--index HCPI]
     [--force] [--validate-only]
"""
from __future__ import annotations

import argparse
import csv
import sys
import time
from collections import Counter
from pathlib import Path

import openpyxl
import pandas as pd

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT))

from Code.common.baselines import run_baselines  # noqa: E402
from Code.common.data import (  # noqa: E402
    FINAL_DATA_PATH,
    FREQ_ALIAS,
    FREQ_TO_SEASONAL_PERIOD,
    build_fixed_window_tsdf,
    load_final_data,
    parse_plan_date,
)
from Code.common.results_io import (  # noqa: E402
    RESULTS_DIR,
    already_has_result,
    write_block_header,
    write_dataframe_block,
    write_run_config,
)

TAB_NAME = "Blk5_Baselines"
BLOCK = "4.1"
EXP_ID_PREFIX = "BL-"
EXPECTED_ROWS = 207
EXPECTED_INDICES = {"HCPI", "FCPI", "CCPI"}
EXPECTED_ORIGINS = list(range(1, 70))
EXPECTED_REGIMES = {"R1": 12, "R2": 12, "R3": 12, "R4": 12, "R5": 12, "R6": 9}
TOTAL_WINDOW = 72
N_PERIODS = 12
METHODS = ["naive", "drift", "atkeson_ohanian", "seasonal_naive", "arima"]

EXPERIMENTS_CSV = REPO_ROOT / "Code" / "experiments.csv"
WORKBOOK_PATH = RESULTS_DIR / f"Results_{TAB_NAME}.xlsx"


def load_block_rows() -> list[dict]:
    with open(EXPERIMENTS_CSV, encoding="utf-8") as f:
        rows = [r for r in csv.DictReader(f)
                if r["block"] == BLOCK and r["exp_id"].startswith(EXP_ID_PREFIX)]
    if len(rows) != EXPECTED_ROWS:
        raise SystemExit(f"Expected exactly {EXPECTED_ROWS} rows for block {BLOCK!r} with exp_id "
                         f"prefix {EXP_ID_PREFIX!r}, got {len(rows)} - filter is wrong, stopping.")
    return rows


def validate_rows(rows: list[dict]) -> None:
    """Asserts the design of `rows` (one or more whole indices) and every row's window
    against Data/final_data.csv. Raises SystemExit listing every failure.
    """
    errors = []
    ids = [r["exp_id"] for r in rows]
    if len(set(ids)) != len(ids):
        errors.append("duplicate exp_ids")
    bad_ids = [i for i in ids if not i.startswith(EXP_ID_PREFIX) or i.startswith(("B21-", "B31-", "B41-", "B42-"))]
    if bad_ids:
        errors.append(f"out-of-scope exp_ids: {bad_ids}")
    bad_block = [r["exp_id"] for r in rows if r["block"] != BLOCK]
    if bad_block:
        errors.append(f"rows outside block {BLOCK}: {bad_block}")

    for index in sorted({r["index"] for r in rows}):
        if index not in EXPECTED_INDICES:
            errors.append(f"unexpected index {index!r}")
        sub = [r for r in rows if r["index"] == index]
        origins = sorted(int(r["origin_no"]) for r in sub)
        if origins != EXPECTED_ORIGINS:
            errors.append(f"{index}: {len(sub)} rows, origins not exactly 1-69 (dupes/gaps): {origins}")
        regimes = Counter(r["regime_of_horizon_start"].split()[0] for r in sub)
        if dict(regimes) != EXPECTED_REGIMES:
            errors.append(f"{index}: regime counts {dict(sorted(regimes.items()))}, expected {EXPECTED_REGIMES}")

    data = pd.read_csv(FINAL_DATA_PATH, parse_dates=["date"])
    data = data[data["frequency"] == "monthly"]
    for r in rows:
        eid = r["exp_id"]
        if r["frequency"] != "Monthly" or int(r["n_periods"]) != N_PERIODS \
                or int(r["total_window"]) != TOTAL_WINDOW or int(r["train_periods"]) != 48 \
                or int(r["validation_periods"]) != 24:
            errors.append(f"{eid}: controls differ (frequency={r['frequency']}, n_periods={r['n_periods']}, "
                          f"train/val/total={r['train_periods']}/{r['validation_periods']}/{r['total_window']})")
        ts, te = parse_plan_date(r["train_data_start"]), parse_plan_date(r["train_data_end"])
        hs, he = parse_plan_date(r["horizon_start"]), parse_plan_date(r["horizon_end"])
        series = data[data["index"] == r["index"]].set_index("date")["value"].dropna()
        n_train = int(((series.index >= ts) & (series.index <= te)).sum())
        n_actual = int(((series.index >= hs) & (series.index <= he)).sum())
        if n_train != TOTAL_WINDOW:
            errors.append(f"{eid}: {n_train} monthly rows {ts:%b %Y}-{te:%b %Y}, expected {TOTAL_WINDOW}")
        if n_actual != N_PERIODS:
            errors.append(f"{eid}: {n_actual} actuals {hs:%b %Y}-{he:%b %Y}, expected {N_PERIODS}")
        if te + pd.DateOffset(months=1) != hs:
            errors.append(f"{eid}: train_data_end {te:%b %Y} is not the month before horizon_start {hs:%b %Y}")

    if errors:
        raise SystemExit("Row validation FAILED - nothing run:\n  " + "\n  ".join(errors))


def print_rows(rows: list[dict]) -> None:
    print(f"{'exp_id':<9} {'index':<5} {'o':>3} {'regime':<4} {'train':<19} {'horizon':<19}")
    for r in rows:
        print(f"{r['exp_id']:<9} {r['index']:<5} {r['origin_no']:>3} {r['regime_of_horizon_start'].split()[0]:<4} "
              f"{r['train_data_start'] + '-' + r['train_data_end']:<19} "
              f"{r['horizon_start'] + '-' + r['horizon_end']:<19}")


def run_one(exp_row: dict) -> float:
    index = exp_row["index"]
    frequency = exp_row["frequency"]
    freq_key = FREQ_ALIAS[frequency]
    seasonal_period = FREQ_TO_SEASONAL_PERIOD[freq_key]

    train_data_start = parse_plan_date(exp_row["train_data_start"])
    train_data_end = parse_plan_date(exp_row["train_data_end"])
    horizon_start = parse_plan_date(exp_row["horizon_start"])
    horizon_end = parse_plan_date(exp_row["horizon_end"])
    prediction_length = int(exp_row["n_periods"])
    total_window = int(exp_row["total_window"])

    series = load_final_data(index, freq_key)
    train_tsdf = build_fixed_window_tsdf(
        series, train_data_start, train_data_end, item_id=index, expected_periods=total_window
    )
    horizon = series[(series["date"] >= horizon_start) & (series["date"] <= horizon_end)]
    if len(horizon) != prediction_length:
        raise ValueError(f"{exp_row['exp_id']}: {len(horizon)} actuals in horizon, expected {prediction_length}")

    past_values = train_tsdf["target"].to_numpy()

    t0 = time.monotonic()
    out = run_baselines(past_values, prediction_length, seasonal_period)
    total_runtime = time.monotonic() - t0

    forecast_df = pd.DataFrame({"date": horizon["date"].to_numpy(), "actual": horizon["value"].to_numpy()})
    for m in METHODS:
        forecast_df[m] = out[f"{m}_forecast"]

    metadata = {
        "exp_id": exp_row["exp_id"], "tab_name": TAB_NAME, "block": BLOCK, "index": index,
        "frequency": frequency, "origin_no": int(exp_row["origin_no"]),
        "regime_of_horizon_start": exp_row["regime_of_horizon_start"],
        "train_data_start": train_data_start.date().isoformat(),
        "train_data_end": train_data_end.date().isoformat(),
        "horizon_start": horizon_start.date().isoformat(),
        "horizon_end": horizon_end.date().isoformat(),
        "n_periods": prediction_length, "total_window": total_window,
        "train_rows_actual": len(past_values), "seasonal_period": seasonal_period,
        **{f"{m}_runtime_sec": round(out[f"{m}_runtime_sec"], 4) for m in METHODS},
        "arima_order": str(out["arima_order"]),
        "total_runtime_sec": round(total_runtime, 4),
    }

    _write_sheet(exp_row["exp_id"], metadata, forecast_df)
    write_run_config(TAB_NAME, exp_row["exp_id"], {
        **metadata, "methods": METHODS,
        "method_notes": "naive=random walk; drift=RW + trailing 12m slope; atkeson_ohanian=last 12m "
                        "inflation carried forward; seasonal_naive=value 12m earlier; arima=grid "
                        "p,d,q in {0,1,2} excl. p=q=0 by AIC (see Code/common/baselines.py)",
    })
    print(f"  {exp_row['exp_id']} ({index}, o{exp_row['origin_no']}): done in {total_runtime:.2f}s "
          f"(ARIMA order {out['arima_order']})", flush=True)
    return total_runtime


def _write_sheet(exp_id: str, metadata: dict, forecast_df: pd.DataFrame) -> None:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    if WORKBOOK_PATH.exists():
        wb = openpyxl.load_workbook(WORKBOOK_PATH)
        if exp_id in wb.sheetnames:
            del wb[exp_id]
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
    ws = wb.create_sheet(title=exp_id)
    row = write_block_header(ws, 1, f"Experiment {exp_id} (baselines: {', '.join(METHODS)})")
    meta_df = pd.DataFrame({"field": list(metadata.keys()), "value": list(metadata.values())})
    row = write_dataframe_block(ws, row, meta_df)
    row = write_block_header(ws, row, "Forecast (index level; all five baseline methods; no ensemble, no weights)")
    write_dataframe_block(ws, row, forecast_df)
    wb.save(WORKBOOK_PATH)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--exp-id", default=None, help="run only this exp_id")
    parser.add_argument("--index", default=None, help='filter to one index, e.g. "HCPI"')
    parser.add_argument("--force", action="store_true", help="re-run even if a result already exists")
    parser.add_argument("--validate-only", action="store_true",
                        help="print and validate the selected rows, then exit without running")
    args = parser.parse_args()

    rows = load_block_rows()
    if args.index:
        rows = [r for r in rows if r["index"] == args.index]
        if not rows:
            raise SystemExit(f"No rows for index {args.index!r} in block {BLOCK}")
    validate_rows(rows)
    if args.validate_only:
        print_rows(rows)
        print(f"\nValidation passed: {len(rows)} rows.")
        return
    if args.exp_id:
        rows = [r for r in rows if r["exp_id"] == args.exp_id]
        if not rows:
            raise SystemExit(f"exp_id {args.exp_id!r} not found in block {BLOCK} (after --index filter)")

    n_run, n_skipped, failures = 0, 0, []
    t_start = time.monotonic()
    for row in rows:
        if not args.force and already_has_result(WORKBOOK_PATH, row["exp_id"]):
            n_skipped += 1
            continue
        try:
            run_one(row)
            n_run += 1
        except Exception as exc:  # one bad origin must not hide the rest - report all at the end
            failures.append((row["exp_id"], repr(exc)))
            print(f"  {row['exp_id']}: FAILED {exc!r}", flush=True)

    print(f"\n{TAB_NAME}: {n_run} run, {n_skipped} skipped, {len(failures)} failed "
          f"(of {len(rows)} selected) in {time.monotonic() - t_start:.1f}s")
    for exp_id, err in failures:
        print(f"  FAILED {exp_id}: {err}")
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
